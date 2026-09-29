"""In-memory export helpers."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document

from queryguard.models import DocumentQueryResponse


def dataframe_to_csv_bytes(frame: pd.DataFrame) -> bytes:
    return frame.to_csv(index=False).encode("utf-8")


def dataframe_to_xlsx_bytes(frame: pd.DataFrame) -> bytes:
    output = BytesIO()
    safe = frame.copy()
    for column in safe.select_dtypes(include="object").columns:
        safe[column] = safe[column].map(
            lambda value: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
        )
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe.to_excel(writer, index=False, sheet_name="result")
    return output.getvalue()


def document_answer_markdown(response: DocumentQueryResponse) -> str:
    lines = ["# Document answer\n", f"**Question:** {response.question}\n", response.answer or ""]
    if response.sources:
        lines.append("\n## Sources")
        for source in response.sources:
            lines.append(f"- **{source.source_name} — {source.locator}:** {source.excerpt}")
    return "\n".join(lines)


def document_answer_docx_bytes(response: DocumentQueryResponse) -> bytes:
    document = Document()
    document.add_heading("Document answer", level=1)
    document.add_paragraph(f"Question: {response.question}")
    document.add_paragraph(response.answer or "")
    if response.sources:
        document.add_heading("Sources", level=2)
        for source in response.sources:
            document.add_paragraph(
                f"{source.source_name} — {source.locator}: {source.excerpt}", style="List Bullet"
            )
    output = BytesIO()
    document.save(output)
    return output.getvalue()
