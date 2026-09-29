"""File-ingestion helpers for databases, spreadsheets, and documents."""

from __future__ import annotations

import re
import sqlite3
import zipfile
from contextlib import closing
from pathlib import Path

import fitz
import pandas as pd
from docx import Document
from pptx import Presentation

from queryguard.database import extract_schema
from queryguard.retrieval import TextChunk

SQLITE_EXTENSIONS = {".db", ".sqlite", ".sqlite3"}
SPREADSHEET_EXTENSIONS = {".csv", ".xlsx"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx"}
INVOICE_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".xlsx", ".csv"}
SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
TABLE_NAME_RE = re.compile(r"[^A-Za-z0-9_]+")


class IngestionError(RuntimeError):
    pass


def safe_filename(name: str) -> str:
    clean = SAFE_NAME.sub("_", Path(name).name).strip("._")
    if not clean:
        raise IngestionError("The uploaded filename is not usable.")
    return clean[:160]


def validate_extension(path: Path, allowed: set[str]) -> None:
    if path.suffix.lower() not in allowed:
        allowed_text = ", ".join(sorted(allowed))
        raise IngestionError(f"Unsupported file type. Allowed extensions: {allowed_text}")


def validate_office_archive(path: Path, max_uncompressed_bytes: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise IngestionError(f"Office file is not a valid ZIP-based document: {path.name}") from exc
    if total > max_uncompressed_bytes:
        raise IngestionError("Office file expands beyond the configured safety limit.")


def validate_sqlite_database(path: Path) -> None:
    validate_extension(path, SQLITE_EXTENSIONS)
    try:
        connection = sqlite3.connect(path)
        try:
            check = connection.execute("PRAGMA quick_check").fetchone()
        finally:
            connection.close()
    except sqlite3.DatabaseError as exc:
        raise IngestionError(f"Invalid SQLite database: {exc}") from exc
    if not check or check[0] != "ok":
        raise IngestionError("SQLite integrity check failed.")
    if not extract_schema(path):
        raise IngestionError("The SQLite database does not contain user tables.")


def safe_table_name(value: str, fallback: str = "data") -> str:
    name = TABLE_NAME_RE.sub("_", value.strip()).strip("_") or fallback
    if name[0].isdigit():
        name = f"table_{name}"
    return name[:60]


def _unique_name(base: str, used: set[str]) -> str:
    candidate = base
    counter = 2
    while candidate.lower() in used:
        candidate = f"{base}_{counter}"
        counter += 1
    used.add(candidate.lower())
    return candidate


def _prepare_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    used: set[str] = set()
    columns: list[str] = []
    for index, column in enumerate(frame.columns, start=1):
        base = safe_table_name(str(column), f"column_{index}")
        columns.append(_unique_name(base, used))
    frame.columns = columns
    return frame


def spreadsheet_to_sqlite(source_path: Path, database_path: Path, max_office_uncompressed_bytes: int) -> list[str]:
    validate_extension(source_path, SPREADSHEET_EXTENSIONS)
    database_path.parent.mkdir(parents=True, exist_ok=True)
    if database_path.exists():
        database_path.unlink()
    warnings: list[str] = []
    tables: list[tuple[str, pd.DataFrame]] = []
    used_names: set[str] = set()
    if source_path.suffix.lower() == ".csv":
        try:
            frame = pd.read_csv(source_path)
        except Exception as exc:
            raise IngestionError(f"Could not read CSV file: {exc}") from exc
        tables.append((_unique_name(safe_table_name(source_path.stem), used_names), _prepare_frame(frame)))
    else:
        validate_office_archive(source_path, max_office_uncompressed_bytes)
        try:
            workbook = pd.ExcelFile(source_path, engine="openpyxl")
        except Exception as exc:
            raise IngestionError(f"Could not read Excel workbook: {exc}") from exc
        for sheet_name in workbook.sheet_names:
            frame = pd.read_excel(workbook, sheet_name=sheet_name)
            if frame.empty and len(frame.columns) == 0:
                warnings.append(f"Skipped empty sheet: {sheet_name}")
                continue
            tables.append((_unique_name(safe_table_name(sheet_name, "sheet"), used_names), _prepare_frame(frame)))
    if not tables:
        raise IngestionError("The spreadsheet did not contain usable tables.")
    try:
        with closing(sqlite3.connect(database_path)) as connection:
            for table_name, frame in tables:
                frame.to_sql(table_name, connection, if_exists="replace", index=False)
    except Exception as exc:
        raise IngestionError(f"Could not convert spreadsheet to SQLite: {exc}") from exc
    return warnings


def _ocr_pdf_page(page: fitz.Page) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        return pytesseract.image_to_string(image).strip()
    except Exception:
        return ""


def _pdf_chunks(path: Path) -> tuple[list[TextChunk], list[str]]:
    chunks: list[TextChunk] = []
    warnings: list[str] = []
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise IngestionError(f"Could not open PDF: {exc}") from exc
    try:
        for page_index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if not text:
                text = _ocr_pdf_page(page)
                if text:
                    warnings.append(f"OCR was used for {path.name} page {page_index}.")
            if text:
                chunks.extend(_split_text(path.name, f"page {page_index}", text))
            else:
                warnings.append(
                    f"No extractable text on {path.name} page {page_index}; "
                    "install the OCR extra and Tesseract for scanned pages."
                )
    finally:
        document.close()
    return chunks, warnings


def _docx_chunks(path: Path, max_uncompressed_bytes: int) -> tuple[list[TextChunk], list[str]]:
    validate_office_archive(path, max_uncompressed_bytes)
    try:
        doc = Document(path)
    except Exception as exc:
        raise IngestionError(f"Could not read Word document: {exc}") from exc
    chunks: list[TextChunk] = []
    section = "document"
    buffer: list[str] = []
    def flush() -> None:
        nonlocal buffer
        if buffer:
            chunks.extend(_split_text(path.name, section, "\n".join(buffer)))
            buffer = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and str(paragraph.style.name).lower().startswith("heading"):
            flush()
            section = text
        else:
            buffer.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if rows:
            chunks.extend(_split_text(path.name, f"table {table_index}", "\n".join(rows)))
    flush()
    return chunks, []


def _pptx_chunks(path: Path, max_uncompressed_bytes: int) -> tuple[list[TextChunk], list[str]]:
    validate_office_archive(path, max_uncompressed_bytes)
    try:
        presentation = Presentation(path)
    except Exception as exc:
        raise IngestionError(f"Could not read PowerPoint file: {exc}") from exc
    chunks: list[TextChunk] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text_parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text.strip())
        if text_parts:
            chunks.extend(_split_text(path.name, f"slide {index}", "\n".join(text_parts)))
    return chunks, []


def _split_text(source: str, locator: str, text: str, max_chars: int = 1200, overlap: int = 150) -> list[TextChunk]:
    clean = " ".join(text.split())
    if not clean:
        return []
    chunks: list[TextChunk] = []
    start = 0
    while start < len(clean):
        end = min(len(clean), start + max_chars)
        if end < len(clean):
            boundary = clean.rfind(" ", start, end)
            if boundary > start + max_chars // 2:
                end = boundary
        chunks.append(TextChunk(source, locator, clean[start:end].strip()))
        if end >= len(clean):
            break
        start = max(start + 1, end - overlap)
    return chunks


def extract_document_chunks(path: Path, max_uncompressed_bytes: int) -> tuple[list[TextChunk], list[str]]:
    validate_extension(path, DOCUMENT_EXTENSIONS)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_chunks(path)
    if suffix == ".docx":
        return _docx_chunks(path, max_uncompressed_bytes)
    return _pptx_chunks(path, max_uncompressed_bytes)
