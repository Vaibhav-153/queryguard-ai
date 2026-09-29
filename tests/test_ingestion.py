from pathlib import Path

import pandas as pd
from docx import Document
from pptx import Presentation

from queryguard.database import extract_schema
from queryguard.ingestion import extract_document_chunks, safe_filename, spreadsheet_to_sqlite


def test_safe_filename_removes_path_and_special_characters():
    assert safe_filename("../../sales report (final).csv") == "sales_report_final_.csv"


def test_csv_is_converted_to_sqlite(tmp_path: Path):
    source = tmp_path / "orders.csv"
    pd.DataFrame({"Order ID": [1, 2], "Amount": [10.5, 20.0]}).to_csv(source, index=False)
    database = tmp_path / "workspace.sqlite"
    spreadsheet_to_sqlite(source, database, 10_000_000)
    schema = extract_schema(database)
    assert len(schema) == 1
    assert schema[0].name == "orders"
    assert [column.name for column in schema[0].columns] == ["Order_ID", "Amount"]


def test_docx_text_is_chunked(tmp_path: Path):
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("Safety", level=1)
    document.add_paragraph("Only validated read-only SQL is executed.")
    document.save(path)
    chunks, warnings = extract_document_chunks(path, 10_000_000)
    assert not warnings
    assert any("read-only SQL" in chunk.text for chunk in chunks)


def test_pptx_text_is_chunked(tmp_path: Path):
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "QueryGuard"
    slide.placeholders[1].text = "Document questions include retrieved evidence."
    presentation.save(path)
    chunks, warnings = extract_document_chunks(path, 10_000_000)
    assert not warnings
    assert any("retrieved evidence" in chunk.text for chunk in chunks)
