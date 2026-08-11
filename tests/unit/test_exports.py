from io import BytesIO

import pandas as pd
from docx import Document
from openpyxl import load_workbook

from queryguard.export.reports import document_answer_docx_bytes, document_answer_markdown
from queryguard.export.tabular import dataframe_to_csv_bytes, dataframe_to_xlsx_bytes
from queryguard.models import DocumentQueryResponse, DocumentSource


def test_csv_export_escapes_formula_like_text():
    frame = pd.DataFrame({"value": ["=2+2", "normal", -5]})
    text = dataframe_to_csv_bytes(frame).decode("utf-8")

    assert "'=2+2" in text
    assert "normal" in text
    assert "-5" in text


def test_xlsx_export_escapes_formula_like_text():
    frame = pd.DataFrame({"value": ["=2+2"]})
    workbook = load_workbook(BytesIO(dataframe_to_xlsx_bytes(frame)), data_only=False)

    assert workbook["Results"]["A2"].value == "'=2+2"


def _document_response() -> DocumentQueryResponse:
    return DocumentQueryResponse(
        status="success",
        question="What is the return window?",
        answer="The return window is 30 days.",
        sources=[
            DocumentSource(
                source_name="policy.pdf",
                locator="Page 2",
                excerpt="Items may be returned within 30 days.",
                score=1.25,
            )
        ],
    )


def test_markdown_report_contains_question_answer_and_evidence():
    text = document_answer_markdown(_document_response())

    assert "What is the return window?" in text
    assert "The return window is 30 days." in text
    assert "policy.pdf" in text
    assert "Page 2" in text


def test_docx_report_can_be_opened_again():
    payload = document_answer_docx_bytes(_document_response())
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "QueryGuard Document Analysis" in text
    assert "What is the return window?" in text
    assert "policy.pdf" in text
