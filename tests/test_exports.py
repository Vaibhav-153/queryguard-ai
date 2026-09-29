from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from queryguard.exports import dataframe_to_xlsx_bytes


def test_excel_export_escapes_formula_like_text():
    frame = pd.DataFrame({"value": ["=2+2", "+cmd", "normal"]})
    payload = dataframe_to_xlsx_bytes(frame)
    workbook = load_workbook(BytesIO(payload), data_only=False)
    values = [workbook["result"].cell(row=i, column=1).value for i in range(2, 5)]
    assert values == ["'=2+2", "'+cmd", "normal"]
